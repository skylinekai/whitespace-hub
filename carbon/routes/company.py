import io
import json
from datetime import date

from flask import (Blueprint, render_template, request, redirect, url_for, flash, g, send_file, abort,
                   session)

import calc
from constants import CATEGORIES, ACTION_STATUSES
from db import q, ex, get_db
from helpers import login_required, current_period, fnum, factor_snapshot
import importer

bp = Blueprint("company", __name__)


def company_factors():
    return q("""SELECT * FROM factors WHERE active=1 AND module IN ('company','both')
                ORDER BY kind, name""")


def need_period():
    p = current_period()
    if not p:
        flash("Create your first reporting period to start adding data.", "info")
    return p


# --------------------------------------------------------------------------
# Dashboard
# --------------------------------------------------------------------------

@bp.route("/")
@login_required()
def dashboard():
    p = current_period()
    if not p:
        return render_template("company/welcome.html")
    s = calc.summarise_period(p)
    prev = q("SELECT * FROM periods WHERE start_date < ? ORDER BY start_date DESC LIMIT 1",
             (p["start_date"],), one=True)
    ps = calc.summarise_period(prev) if prev else None
    target = q("SELECT * FROM targets ORDER BY id LIMIT 1", one=True)
    traj = trajectory(target) if target else None
    recs = calc.company_recommendations(s)
    chart = dict(
        scopes=[s["by_scope"][1] / 1000, s["by_scope"][2] / 1000, s["by_scope"][3] / 1000],
        cats=[[CATEGORIES[c][1], round(kg / 1000, 3), CATEGORIES[c][0]] for c, kg in s["cat_sorted"]],
        months=list(s["by_month"].keys()), month_vals=[round(v / 1000, 3) for v in s["by_month"].values()],
        traj=traj,
    )
    return render_template("company/dashboard.html", s=s, ps=ps, prev=prev, chart=chart, recs=recs,
                           target=target, traj=traj)


def trajectory(target):
    """Linear path from base-period emissions to target year, plus actuals per period."""
    base_p = q("SELECT * FROM periods WHERE id=?", (target["base_period_id"],), one=True)
    if not base_p:
        return None
    base_s = calc.summarise_period(base_p)
    base_kg = calc.coverage_total(base_s, target["coverage"])
    base_year = int(base_p["end_date"][:4])
    ty = int(target["target_year"])
    years = list(range(base_year, max(ty, base_year + 1) + 1))
    path = [base_kg * (1 - target["reduction_pct"] / 100 * (y - base_year) / max(ty - base_year, 1)) / 1000
            for y in years]
    actual = {}
    for p in q("SELECT * FROM periods ORDER BY start_date"):
        y = int(p["end_date"][:4])
        if y in years:
            actual[y] = calc.coverage_total(calc.summarise_period(p), target["coverage"]) / 1000
    planned = sum((a["est_reduction_t"] or 0) for a in q(
        "SELECT est_reduction_t FROM actions WHERE status IN ('planned','in_progress','done')"))
    latest = actual[max(actual)] if actual else base_kg / 1000
    return dict(years=years, path=[round(v, 3) for v in path],
                actual=[round(actual[y], 3) if y in actual else None for y in years],
                base_t=base_kg / 1000, target_t=base_kg / 1000 * (1 - target["reduction_pct"] / 100),
                base_year=base_year, planned_t=planned, latest_t=latest,
                with_actions_t=max(latest - planned, 0))


# --------------------------------------------------------------------------
# Activities
# --------------------------------------------------------------------------

@bp.route("/activities")
@login_required()
def activities():
    p = need_period()
    if not p:
        return redirect(url_for("company.periods"))
    sql = """SELECT a.*, s.name AS site_name, u.name AS user_name FROM activities a
             LEFT JOIN sites s ON s.id=a.site_id LEFT JOIN users u ON u.id=a.created_by
             WHERE a.period_id=?"""
    args = [p["id"]]
    cat = request.args.get("category")
    if cat:
        sql += " AND a.category=?"
        args.append(cat)
    site = request.args.get("site")
    if site:
        sql += " AND a.site_id=?"
        args.append(site)
    term = request.args.get("q", "").strip()
    if term:
        sql += " AND (a.description LIKE ? OR a.ef_name LIKE ? OR a.evidence LIKE ?)"
        args += [f"%{term}%"] * 3
    sql += " ORDER BY a.activity_date DESC, a.id DESC"
    rows = []
    for a in q(sql, args):
        d, u, m = calc.activity_parts(a)
        rows.append(dict(a=a, direct=d, upstream=u))
    total = sum(r["direct"] + r["upstream"] for r in rows)
    return render_template("company/activities.html", rows=rows, total=total, period=p,
                           sites=q("SELECT * FROM sites ORDER BY name"))


def save_activity(aid, p):
    f = q("SELECT * FROM factors WHERE id=?", (request.form.get("factor_id"),), one=True)
    qty = fnum(request.form.get("quantity"))
    if not f or qty is None:
        flash("Choose an emission factor and enter a quantity.", "error")
        return False
    d = request.form.get("activity_date") or None
    if d and not (p["start_date"] <= d <= p["end_date"]):
        flash(f"That date is outside {p['name']} ({p['start_date']} to {p['end_date']}).", "error")
        return False
    vals = dict(period_id=p["id"], site_id=request.form.get("site_id") or None,
                category=request.form.get("category") or f["default_category"],
                factor_id=f["id"], description=request.form.get("description", "").strip(),
                activity_date=d, quantity=qty, unit=f["unit"],
                renewable=1 if request.form.get("renewable") else 0,
                data_quality=request.form.get("data_quality", "actual"),
                evidence=request.form.get("evidence", "").strip(), **factor_snapshot(f))
    if aid:
        sets = ",".join(f"{k}=:{k}" for k in vals)
        ex(f"UPDATE activities SET {sets} WHERE id=:id", dict(vals, id=aid))
    else:
        vals["created_by"] = g.user["id"]
        cols = ",".join(vals)
        ex(f"INSERT INTO activities({cols}) VALUES({','.join(':' + k for k in vals)})", vals)
    return True


@bp.route("/activities/new", methods=["GET", "POST"])
@bp.route("/activities/<int:aid>/edit", methods=["GET", "POST"])
@login_required("editor")
def activity_form(aid=None):
    p = need_period()
    if not p:
        return redirect(url_for("company.periods"))
    a = q("SELECT * FROM activities WHERE id=?", (aid,), one=True) if aid else None
    if aid and not a:
        abort(404)
    if p["locked"]:
        flash(f"{p['name']} is locked. Unlock it under Reporting periods to change data.", "error")
        return redirect(url_for("company.activities"))
    if request.method == "POST":
        if save_activity(aid, p):
            flash("Saved.", "ok")
            if request.form.get("again"):
                return redirect(url_for("company.activity_form"))
            return redirect(url_for("company.activities"))
    return render_template("company/activity_form.html", a=a, period=p, factors=company_factors(),
                           sites=q("SELECT * FROM sites WHERE active=1 ORDER BY name"),
                           factors_json=json.dumps({f["id"]: dict(unit=f["unit"], cat=f["default_category"],
                                                                  kind=f["kind"], kg=f["kg_per_unit"],
                                                                  up=f["wtt_per_unit"] + f["td_per_unit"])
                                                    for f in company_factors()}))


@bp.route("/activities/<int:aid>/delete", methods=["POST"])
@login_required("editor")
def activity_delete(aid):
    a = q("SELECT a.*, p.locked FROM activities a JOIN periods p ON p.id=a.period_id WHERE a.id=?", (aid,), one=True)
    if a and not a["locked"]:
        ex("DELETE FROM activities WHERE id=?", (aid,))
        flash("Deleted.", "ok")
    return redirect(url_for("company.activities"))


@bp.route("/activities/export.xlsx")
@login_required()
def export_activities():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    p = current_period()
    wb = Workbook()
    ws = wb.active
    ws.title = "Activities"
    hdr = ["Date", "Site", "Scope", "Category", "Description", "Factor code", "Factor", "Quantity", "Unit",
           "Factor kgCO2e/unit", "Direct kgCO2e", "Upstream kgCO2e (WTT/T&D)", "Total kgCO2e",
           "Renewable tariff", "Data quality", "Evidence", "Factor source"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="24215C")
    for a in q("""SELECT a.*, s.name AS site_name FROM activities a LEFT JOIN sites s ON s.id=a.site_id
                  WHERE period_id=? ORDER BY activity_date""", (p["id"],)):
        d, u, _ = calc.activity_parts(a)
        ws.append([a["activity_date"], a["site_name"], CATEGORIES[a["category"]][0], CATEGORIES[a["category"]][1],
                   a["description"], a["ef_code"], a["ef_name"], a["quantity"], a["unit"], a["ef_kg"],
                   round(d, 3), round(u, 3), round(d + u, 3), "Yes" if a["renewable"] else "",
                   a["data_quality"], a["evidence"], a["ef_source"]])
    for col, w in zip("ABCDEFGHIJKLMNOPQ", [11, 18, 7, 32, 34, 20, 34, 11, 11, 12, 13, 16, 13, 10, 12, 24, 18]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"activities-{p['name']}.xlsx")


# --------------------------------------------------------------------------
# Import (spreadsheet + AI)
# --------------------------------------------------------------------------

@bp.route("/import", methods=["GET", "POST"])
@login_required("editor")
def import_data():
    p = need_period()
    if not p:
        return redirect(url_for("company.periods"))
    if request.method == "POST":
        up = request.files.get("file")
        if not up or not up.filename:
            flash("Choose a file to upload.", "error")
            return redirect(url_for("company.import_data"))
        data = up.read()
        mode = request.form.get("mode", "template")
        factors = company_factors()
        sites = q("SELECT * FROM sites WHERE active=1")
        try:
            if mode == "ai":
                raw = importer.ai_extract(up.filename, data, factors, [s["name"] for s in sites])
            else:
                raw = importer.read_table(up.filename, data)
        except Exception as e:  # noqa: BLE001 — show any parsing/API problem to the user
            flash(f"Couldn't read that file: {e}", "error")
            return redirect(url_for("company.import_data"))
        rows = importer.validate_rows(raw, {f["code"]: f for f in factors},
                                      {s["name"].lower(): s for s in sites}, p)
        if not rows:
            flash("No rows found in that file.", "error")
            return redirect(url_for("company.import_data"))
        return render_template("company/import_review.html", rows=rows, period=p, mode=mode,
                               filename=up.filename, rows_json=json.dumps(rows))
    return render_template("company/import.html", period=p)


@bp.route("/import/confirm", methods=["POST"])
@login_required("editor")
def import_confirm():
    p = current_period()
    rows = json.loads(request.form.get("rows_json", "[]"))
    keep = set(request.form.getlist("keep"))
    added = 0
    db = get_db()
    for r in rows:
        if str(r["row"]) not in keep or not r["ok"]:
            continue
        f = q("SELECT * FROM factors WHERE code=?", (r["factor_code"],), one=True)
        if not f:
            continue
        cat = r["category"] if r["category"] in CATEGORIES else f["default_category"]
        vals = dict(period_id=p["id"], site_id=r["site_id"] or None, category=cat, factor_id=f["id"],
                    description=r["description"], activity_date=r["date"] or None, quantity=r["quantity"],
                    unit=f["unit"], renewable=1 if r["renewable"] else 0,
                    data_quality=r["data_quality"] if r["data_quality"] in ("actual", "estimated", "spend") else "actual",
                    evidence=r["evidence"], created_by=g.user["id"], **factor_snapshot(f))
        db.execute(f"INSERT INTO activities({','.join(vals)}) VALUES({','.join(':' + k for k in vals)})", vals)
        added += 1
    db.commit()
    flash(f"Imported {added} activit{'y' if added == 1 else 'ies'} into {p['name']}.", "ok")
    return redirect(url_for("company.activities"))


@bp.route("/import/template.xlsx")
@login_required()
def import_template():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Activities"
    ws.append(importer.TEMPLATE_COLUMNS)
    ws.append(["2026-01-31", "Head office", "ELEC-UK-KWH", 4200, "January electricity bill", "", "actual", "no",
               "Octopus inv 12345"])
    ws.append(["2026-01-31", "Workshop", "FUEL-DIESEL-L", 380, "Fuel card — vans", "", "actual", "", "Allstar Jan"])
    ws.append(["2026-02-14", "", "TRV-FLY-SH-ECO-PKM", 2400, "Crew flights to Frankfurt (3 people, return)", "",
               "estimated", "", "Travel booking"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="24215C")
    for col, w in zip("ABCDEFGHI", [12, 16, 22, 10, 40, 22, 12, 10, 22]):
        ws.column_dimensions[col].width = w
    ref = wb.create_sheet("Factor codes")
    ref.append(["factor_code", "name", "unit", "default category", "kgCO2e per unit", "source"])
    for c in ref[1]:
        c.font = Font(bold=True)
    for f in company_factors():
        ref.append([f["code"], f["name"], f["unit"], CATEGORIES[f["default_category"]][1], f["kg_per_unit"], f["source"]])
    for col, w in zip("ABCDEF", [24, 46, 12, 34, 14, 18]):
        ref.column_dimensions[col].width = w
    sites = wb.create_sheet("Sites")
    sites.append(["site (use these names exactly)"])
    for s in q("SELECT name FROM sites WHERE active=1 ORDER BY name"):
        sites.append([s["name"]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="carbon-activity-import-template.xlsx")


# --------------------------------------------------------------------------
# Periods & sites
# --------------------------------------------------------------------------

@bp.route("/periods")
@login_required()
def periods():
    rows = []
    for p in q("SELECT * FROM periods ORDER BY start_date DESC"):
        s = calc.summarise_period(p)
        rows.append(dict(p=p, s=s))
    return render_template("company/periods.html", rows=rows)


@bp.route("/periods/new", methods=["GET", "POST"])
@bp.route("/periods/<int:pid>/edit", methods=["GET", "POST"])
@login_required("editor")
def period_form(pid=None):
    p = q("SELECT * FROM periods WHERE id=?", (pid,), one=True) if pid else None
    if request.method == "POST":
        f = request.form
        if not f.get("name") or not f.get("start_date") or not f.get("end_date") or f["start_date"] > f["end_date"]:
            flash("Give the period a name and a valid start and end date.", "error")
        else:
            vals = (f["name"].strip(), f["start_date"], f["end_date"], fnum(f.get("revenue_gbp")),
                    fnum(f.get("employees_fte")), 1 if f.get("include_stands") else 0,
                    1 if f.get("locked") else 0, f.get("notes", ""))
            if p:
                ex("""UPDATE periods SET name=?,start_date=?,end_date=?,revenue_gbp=?,employees_fte=?,
                      include_stands=?,locked=?,notes=? WHERE id=?""", vals + (pid,))
            else:
                pid = ex("""INSERT INTO periods(name,start_date,end_date,revenue_gbp,employees_fte,include_stands,
                            locked,notes) VALUES(?,?,?,?,?,?,?,?)""", vals)
                session["period_id"] = pid
            flash("Reporting period saved.", "ok")
            return redirect(url_for("company.periods"))
    y = date.today().year
    return render_template("company/period_form.html", p=p, default_start=f"{y}-01-01", default_end=f"{y}-12-31")


@bp.route("/periods/<int:pid>/delete", methods=["POST"])
@login_required("admin")
def period_delete(pid):
    ex("DELETE FROM periods WHERE id=?", (pid,))
    session.pop("period_id", None)
    flash("Reporting period and its activity data deleted.", "ok")
    return redirect(url_for("company.periods"))


@bp.route("/periods/<int:pid>/refresh-factors", methods=["POST"])
@login_required("admin")
def period_refresh(pid):
    p = q("SELECT * FROM periods WHERE id=?", (pid,), one=True)
    if p["locked"]:
        flash("Unlock the period before re-applying factors.", "error")
        return redirect(url_for("company.periods"))
    n = 0
    db = get_db()
    for a in q("SELECT a.id, f.* FROM activities a JOIN factors f ON f.id=a.factor_id WHERE a.period_id=?", (pid,)):
        snap = factor_snapshot(a)
        db.execute("""UPDATE activities SET ef_code=:ef_code, ef_name=:ef_name, ef_kg=:ef_kg, ef_wtt=:ef_wtt,
                      ef_td=:ef_td, ef_kwh=:ef_kwh, ef_source=:ef_source WHERE id=:id""", dict(snap, id=a[0]))
        n += 1
    db.commit()
    flash(f"Re-applied current emission factors to {n} activities in {p['name']}.", "ok")
    return redirect(url_for("company.periods"))


@bp.route("/sites", methods=["GET", "POST"])
@login_required()
def sites():
    if request.method == "POST":
        if g.user["role"] == "viewer":
            abort(403)
        f = request.form
        vals = (f.get("name", "").strip(), f.get("site_type", ""), f.get("address", ""),
                f.get("country", "United Kingdom"), fnum(f.get("floor_area_m2")), 1 if f.get("active") else 0)
        if not vals[0]:
            flash("Site name is required.", "error")
        elif f.get("id"):
            ex("UPDATE sites SET name=?,site_type=?,address=?,country=?,floor_area_m2=?,active=? WHERE id=?",
               vals + (f["id"],))
            flash("Site updated.", "ok")
        else:
            ex("INSERT INTO sites(name,site_type,address,country,floor_area_m2,active) VALUES(?,?,?,?,?,?)", vals)
            flash("Site added.", "ok")
        return redirect(url_for("company.sites"))
    edit = q("SELECT * FROM sites WHERE id=?", (request.args.get("edit"),), one=True) if request.args.get("edit") else None
    return render_template("company/sites.html", sites=q("SELECT * FROM sites ORDER BY active DESC, name"), edit=edit)


# --------------------------------------------------------------------------
# Targets & action plan
# --------------------------------------------------------------------------

@bp.route("/targets")
@login_required()
def targets():
    target = q("SELECT t.*, p.name AS base_name FROM targets t LEFT JOIN periods p ON p.id=t.base_period_id "
               "ORDER BY t.id LIMIT 1", one=True)
    traj = trajectory(target) if target else None
    acts = q("""SELECT * FROM actions ORDER BY CASE status WHEN 'in_progress' THEN 0 WHEN 'planned' THEN 1
                WHEN 'idea' THEN 2 ELSE 3 END, due_date""")
    p = current_period()
    recs = calc.company_recommendations(calc.summarise_period(p)) if p else []
    return render_template("company/targets.html", target=target, traj=traj, actions=acts, recs=recs,
                           periods=q("SELECT * FROM periods ORDER BY start_date"), statuses=ACTION_STATUSES)


@bp.route("/targets/save", methods=["POST"])
@login_required("editor")
def target_save():
    f = request.form
    pct = fnum(f.get("reduction_pct"))
    ty = fnum(f.get("target_year"))
    if not f.get("base_period_id") or pct is None or ty is None:
        flash("Pick a base period, target year and reduction %.", "error")
        return redirect(url_for("company.targets"))
    vals = (f.get("name") or "Near-term reduction target", f["base_period_id"], int(ty), pct,
            f.get("coverage", "s12"), f.get("notes", ""))
    if f.get("id"):
        ex("UPDATE targets SET name=?,base_period_id=?,target_year=?,reduction_pct=?,coverage=?,notes=? WHERE id=?",
           vals + (f["id"],))
    else:
        ex("INSERT INTO targets(name,base_period_id,target_year,reduction_pct,coverage,notes) VALUES(?,?,?,?,?,?)", vals)
    flash("Target saved.", "ok")
    return redirect(url_for("company.targets"))


@bp.route("/actions/save", methods=["POST"])
@login_required("editor")
def action_save():
    f = request.form
    if not f.get("title"):
        flash("An action needs a title.", "error")
        return redirect(url_for("company.targets"))
    vals = (f["title"].strip(), f.get("description", ""), f.get("category", ""), fnum(f.get("est_reduction_t")),
            fnum(f.get("cost_gbp")), f.get("owner", ""), f.get("status", "idea"), f.get("due_date") or None)
    if f.get("id"):
        ex("""UPDATE actions SET title=?,description=?,category=?,est_reduction_t=?,cost_gbp=?,owner=?,status=?,
              due_date=? WHERE id=?""", vals + (f["id"],))
    else:
        ex("""INSERT INTO actions(title,description,category,est_reduction_t,cost_gbp,owner,status,due_date)
              VALUES(?,?,?,?,?,?,?,?)""", vals)
    flash("Action saved.", "ok")
    return redirect(url_for("company.targets") + "#actions")


@bp.route("/actions/<int:aid>/delete", methods=["POST"])
@login_required("editor")
def action_delete(aid):
    ex("DELETE FROM actions WHERE id=?", (aid,))
    return redirect(url_for("company.targets") + "#actions")
