"""Spreadsheet / CSV import and optional AI extraction of activity data."""
import base64
import csv
import io
import json
import os
import re
from datetime import date, datetime

import requests

TEMPLATE_COLUMNS = ["date", "site", "factor_code", "quantity", "description", "category",
                    "data_quality", "renewable", "evidence"]


def _norm(h):
    return re.sub(r"[^a-z_]", "", str(h or "").strip().lower().replace(" ", "_"))


def _cell(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v).strip()


def read_table(filename, data):
    """Return list of dicts from a CSV or XLSX upload."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_norm(h) for h in rows[0]]
        return [{headers[i]: _cell(v) for i, v in enumerate(r) if i < len(headers)} for r in rows[1:]
                if any(c not in (None, "") for c in r)]
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{_norm(k): (v or "").strip() for k, v in r.items()} for r in reader
            if any((v or "").strip() for v in r.values())]


def parse_date(s):
    s = (s or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%Y/%m/%d", "%b %Y", "%B %Y", "%Y-%m"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def validate_rows(raw_rows, factors_by_code, sites_by_name, period):
    """Normalise rows for the review screen. Each row gets an 'errors' list."""
    out = []
    for i, r in enumerate(raw_rows, start=1):
        errs = []
        code = (r.get("factor_code") or "").strip().upper()
        f = factors_by_code.get(code)
        if not f:
            errs.append(f"Unknown factor code '{code}'")
        qty_s = str(r.get("quantity", "")).replace(",", "").replace("£", "").strip()
        try:
            qty = float(qty_s)
        except ValueError:
            qty = None
            errs.append("Quantity is not a number")
        d = parse_date(r.get("date")) if r.get("date") else None
        if r.get("date") and not d:
            errs.append(f"Date '{r.get('date')}' not understood (use YYYY-MM-DD)")
        if d and period and not (period["start_date"] <= d <= period["end_date"]):
            errs.append(f"Date {d} is outside {period['name']}")
        blocking = bool(errs)
        site_name = (r.get("site") or "").strip()
        site = sites_by_name.get(site_name.lower()) if site_name else None
        if site_name and not site:
            errs.append(f"Site '{site_name}' not found (row will be unassigned)")
        cat = (r.get("category") or "").strip() or (f["default_category"] if f else "")
        out.append(dict(
            row=i, factor_code=code, factor_name=f["name"] if f else "", unit=f["unit"] if f else "",
            quantity=qty, date=d or "", site_id=site["id"] if site else "", site=site_name,
            description=r.get("description", ""), category=cat,
            data_quality=(r.get("data_quality") or "actual").strip().lower() or "actual",
            renewable=str(r.get("renewable", "")).strip().lower() in ("1", "yes", "y", "true"),
            evidence=r.get("evidence", ""), confidence=r.get("confidence", ""),
            errors=errs, ok=not blocking,
        ))
    return out


# --------------------------------------------------------------------------
# AI extraction (optional — needs ANTHROPIC_API_KEY)
# --------------------------------------------------------------------------

AI_PROMPT = """You are a carbon accounting data assistant for Skyline Whitespace, a UK exhibition stand
design-and-build company. Extract every greenhouse-gas-relevant activity line from the attached
document (utility bills, fuel cards, invoices, travel bookings, waste notes, mileage claims, supplier invoices).

Map each line to ONE factor code from this list (code | name | unit):
{factors}

Rules:
- quantity must be in the factor's unit. Prefer physical units (kWh, litres, km, nights, tonnes) over spend.
  Use a spend factor (unit GBP, net of VAT) only when no physical quantity is available.
- For electricity and gas bills use the consumption in kWh for the billing period. Use the bill end date.
- Flights: estimate great-circle distance per passenger for the whole trip (return = both legs).
- date in YYYY-MM-DD. If only a month is known use the last day of that month.
- confidence: high / medium / low.
- Do not invent lines that are not in the document.

Reply with ONLY a JSON array, no prose, of objects with keys:
date, factor_code, quantity, description, site, data_quality ("actual" or "estimated" or "spend"),
evidence (supplier name + invoice/bill reference if shown), confidence.
"""


def ai_extract(filename, data, factors, site_names):
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        raise RuntimeError("AI import is not switched on. Add ANTHROPIC_API_KEY in Render to enable it.")
    model = os.environ.get("ANTHROPIC_MODEL", "claude-sonnet-5-5")
    factor_list = "\n".join(f"{f['code']} | {f['name']} | {f['unit']}" for f in factors)
    prompt = AI_PROMPT.format(factors=factor_list)
    if site_names:
        prompt += "\nKnown sites (use the closest match in 'site', else leave blank): " + ", ".join(site_names)

    name = filename.lower()
    content = []
    if name.endswith(".pdf"):
        content.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                        "data": base64.b64encode(data).decode()}})
    elif name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        mt = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp",
              "gif": "image/gif"}[name.rsplit(".", 1)[1]]
        content.append({"type": "image", "source": {"type": "base64", "media_type": mt,
                                                     "data": base64.b64encode(data).decode()}})
    elif name.endswith((".xlsx", ".xlsm")):
        rows = read_table(filename, data)
        content.append({"type": "text", "text": "Spreadsheet rows (JSON):\n" + json.dumps(rows[:800])})
    else:
        content.append({"type": "text", "text": "Document text:\n" + data.decode("utf-8", errors="replace")[:150000]})
    content.append({"type": "text", "text": prompt})

    resp = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
        json={"model": model, "max_tokens": 8000, "messages": [{"role": "user", "content": content}]},
        timeout=180,
    )
    if resp.status_code != 200:
        raise RuntimeError(f"AI service error {resp.status_code}: {resp.text[:300]}")
    text = "".join(b.get("text", "") for b in resp.json().get("content", []) if b.get("type") == "text")
    m = re.search(r"\[.*\]", text, re.S)
    if not m:
        raise RuntimeError("The AI did not return any activity lines for this file.")
    rows = json.loads(m.group(0))
    for r in rows:
        r["evidence"] = (r.get("evidence") or filename)[:200]
        r["quantity"] = str(r.get("quantity", ""))
    return rows
